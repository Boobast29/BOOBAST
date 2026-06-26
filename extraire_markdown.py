#!/usr/bin/env python3
"""Extraction de texte (en Markdown) et d'images depuis des documents.

Petite application avec interface graphique : tu choisis un dossier contenant
tes fichiers PDF, Word (.docx) ou PowerPoint (.pptx), un dossier de sortie, et
l'outil génère pour chaque document :

    <dossier_sortie>/<nom_du_document>/
        <nom_du_document>.md   -> le texte du document, converti en Markdown
        images/                -> les images extraites du document

Lancement :
    python extraire_markdown.py
"""

from __future__ import annotations

import os
import queue
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

EXTENSIONS_SUPPORTEES = {".pdf", ".docx", ".pptx"}


# --- Extraction PDF (texte + images) via PyMuPDF -----------------------------


def extraire_pdf(chemin: Path, dossier_images: Path) -> str:
    import fitz  # PyMuPDF

    morceaux = []
    compteur_image = 0
    document = fitz.open(str(chemin))
    try:
        for numero_page, page in enumerate(document, start=1):
            texte = page.get_text().strip()
            if texte:
                morceaux.append(f"## Page {numero_page}\n\n{texte}")

            for image_info in page.get_images(full=True):
                xref = image_info[0]
                compteur_image += 1
                try:
                    image = document.extract_image(xref)
                except Exception:
                    continue
                nom_image = f"image_{compteur_image:03d}.{image['ext']}"
                (dossier_images / nom_image).write_bytes(image["image"])
                morceaux.append(f"![{nom_image}](images/{nom_image})")
    finally:
        document.close()

    return "\n\n".join(morceaux)


# --- Extraction DOCX (texte + images) ----------------------------------------


def extraire_docx(chemin: Path, dossier_images: Path) -> str:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(chemin))
    morceaux = []
    compteur_image = 0

    # Sauvegarde toutes les images intégrées au document.
    images_par_rid = {}
    for rid, part in document.part.rels.items():
        if "image" in part.reltype:
            compteur_image += 1
            extension = Path(part.target_ref).suffix or ".png"
            nom_image = f"image_{compteur_image:03d}{extension}"
            (dossier_images / nom_image).write_bytes(part.target_part.blob)
            images_par_rid[part.rId if hasattr(part, "rId") else rid] = nom_image

    for paragraphe in document.paragraphs:
        style = (paragraphe.style.name or "").lower()
        texte = paragraphe.text.strip()

        # Repère les images référencées dans ce paragraphe.
        for run in paragraphe.runs:
            for blip in run._element.findall(
                ".//" + qn("a:blip")
            ):
                rid = blip.get(qn("r:embed"))
                nom_image = images_par_rid.get(rid)
                if nom_image:
                    morceaux.append(f"![{nom_image}](images/{nom_image})")

        if not texte:
            continue
        if style.startswith("heading 1") or style == "title":
            morceaux.append(f"# {texte}")
        elif style.startswith("heading 2"):
            morceaux.append(f"## {texte}")
        elif style.startswith("heading"):
            morceaux.append(f"### {texte}")
        else:
            morceaux.append(texte)

    return "\n\n".join(morceaux)


# --- Extraction PPTX (texte + images) ----------------------------------------


def extraire_pptx(chemin: Path, dossier_images: Path) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(chemin))
    morceaux = []
    compteur_image = 0

    for numero_diapo, diapo in enumerate(presentation.slides, start=1):
        morceaux.append(f"## Diapositive {numero_diapo}")
        for forme in diapo.shapes:
            if forme.has_text_frame and forme.text_frame.text.strip():
                morceaux.append(forme.text_frame.text.strip())
            if forme.shape_type == MSO_SHAPE_TYPE.PICTURE:
                compteur_image += 1
                image = forme.image
                nom_image = f"image_{compteur_image:03d}.{image.ext}"
                (dossier_images / nom_image).write_bytes(image.blob)
                morceaux.append(f"![{nom_image}](images/{nom_image})")

    return "\n\n".join(morceaux)


def extraire(chemin: Path, dossier_images: Path) -> str:
    extension = chemin.suffix.lower()
    if extension == ".pdf":
        return extraire_pdf(chemin, dossier_images)
    if extension == ".docx":
        return extraire_docx(chemin, dossier_images)
    if extension == ".pptx":
        return extraire_pptx(chemin, dossier_images)
    raise ValueError(f"Extension non prise en charge : {extension}")


# --- Logique d'extraction (exécutée dans un thread de fond) -----------------


def nom_sortie_unique(dossier: Path, nom: str) -> Path:
    """Évite d'écraser un dossier de sortie portant déjà ce nom."""
    cible = dossier / nom
    if not cible.exists():
        return cible
    compteur = 1
    while True:
        candidat = dossier / f"{nom} ({compteur})"
        if not candidat.exists():
            return candidat
        compteur += 1


def traiter_dossier(dossier_entree: Path, dossier_sortie: Path, journaliser, fini):
    """Parcourt le dossier d'entrée et extrait chaque document. Thread de fond."""
    fichiers = [
        p for p in sorted(dossier_entree.iterdir())
        if p.is_file() and p.suffix.lower() in EXTENSIONS_SUPPORTEES
    ]

    if not fichiers:
        journaliser("Aucun fichier PDF, Word (.docx) ou PowerPoint (.pptx) "
                    "trouvé dans ce dossier.")
        fini()
        return

    journaliser(f"{len(fichiers)} document(s) à traiter.\n")

    for i, fichier in enumerate(fichiers, start=1):
        journaliser(f"[{i}/{len(fichiers)}] {fichier.name} — extraction…")
        try:
            dossier_doc = nom_sortie_unique(dossier_sortie, fichier.stem)
            dossier_doc.mkdir(parents=True)
            dossier_images = dossier_doc / "images"
            dossier_images.mkdir()

            markdown = extraire(fichier, dossier_images)

            if not any(dossier_images.iterdir()):
                dossier_images.rmdir()

            chemin_md = dossier_doc / f"{fichier.stem}.md"
            chemin_md.write_text(markdown, encoding="utf-8")

            journaliser(f"    → {dossier_doc}\n")
        except Exception as erreur:  # noqa: BLE001 - on continue malgré une erreur isolée
            journaliser(f"    ÉCHEC : {erreur}\n")

    journaliser("✅ Extraction terminée.")
    fini()


# --- Interface graphique -----------------------------------------------------


class Application(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Extraction Markdown + images")
        self.geometry("720x560")
        self.minsize(640, 480)

        self.file_journal: queue.Queue[str] = queue.Queue()
        self._construire_interface()
        self.after(100, self._vider_journal)

    def _construire_interface(self):
        cadre = ttk.Frame(self, padding=12)
        cadre.pack(fill="both", expand=True)

        # Dossier d'entrée
        ttk.Label(cadre, text="Dossier contenant tes documents (PDF, .docx, .pptx) :").pack(anchor="w")
        ligne = ttk.Frame(cadre)
        ligne.pack(fill="x", pady=(2, 10))
        self.var_entree = tk.StringVar()
        ttk.Entry(ligne, textvariable=self.var_entree).pack(
            side="left", fill="x", expand=True)
        ttk.Button(ligne, text="Parcourir…", command=self._choisir_entree).pack(
            side="left", padx=(6, 0))

        # Dossier de sortie
        ttk.Label(cadre, text="Dossier de sortie (un sous-dossier par document) :").pack(anchor="w")
        ligne2 = ttk.Frame(cadre)
        ligne2.pack(fill="x", pady=(2, 10))
        self.var_sortie = tk.StringVar()
        ttk.Entry(ligne2, textvariable=self.var_sortie).pack(
            side="left", fill="x", expand=True)
        ttk.Button(ligne2, text="Parcourir…", command=self._choisir_sortie).pack(
            side="left", padx=(6, 0))

        # Bouton lancer
        self.bouton_lancer = ttk.Button(
            cadre, text="Lancer l'extraction", command=self._lancer)
        self.bouton_lancer.pack(anchor="w", pady=(0, 10))

        # Journal
        ttk.Label(cadre, text="Journal :").pack(anchor="w")
        self.journal = scrolledtext.ScrolledText(cadre, height=18, state="disabled")
        self.journal.pack(fill="both", expand=True, pady=(2, 0))

    def _choisir_entree(self):
        dossier = filedialog.askdirectory(title="Choisis le dossier à traiter")
        if dossier:
            self.var_entree.set(dossier)

    def _choisir_sortie(self):
        dossier = filedialog.askdirectory(title="Choisis le dossier de sortie")
        if dossier:
            self.var_sortie.set(dossier)

    def _ecrire(self, message: str):
        self.journal.configure(state="normal")
        self.journal.insert("end", message + "\n")
        self.journal.see("end")
        self.journal.configure(state="disabled")

    def _vider_journal(self):
        while not self.file_journal.empty():
            self._ecrire(self.file_journal.get_nowait())
        self.after(100, self._vider_journal)

    def _lancer(self):
        entree_txt = self.var_entree.get().strip()
        if not entree_txt:
            messagebox.showwarning("Dossier manquant", "Choisis d'abord le dossier à traiter.")
            return
        dossier_entree = Path(entree_txt)
        if not dossier_entree.is_dir():
            messagebox.showerror("Dossier invalide", "Ce dossier n'existe pas.")
            return

        sortie_txt = self.var_sortie.get().strip()
        if not sortie_txt:
            messagebox.showwarning("Dossier manquant", "Choisis d'abord le dossier de sortie.")
            return
        dossier_sortie = Path(sortie_txt)
        dossier_sortie.mkdir(parents=True, exist_ok=True)

        self.bouton_lancer.configure(state="disabled")
        self.journal.configure(state="normal")
        self.journal.delete("1.0", "end")
        self.journal.configure(state="disabled")

        threading.Thread(
            target=traiter_dossier,
            args=(
                dossier_entree,
                dossier_sortie,
                self.file_journal.put,
                lambda: self.after(0, self._reactiver),
            ),
            daemon=True,
        ).start()

    def _reactiver(self):
        self.bouton_lancer.configure(state="normal")


if __name__ == "__main__":
    Application().mainloop()
